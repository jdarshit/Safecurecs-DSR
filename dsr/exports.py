"""Excel export for DSR reports. Business logic (workbook construction) is
kept here, out of views and templates, per project convention."""

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL_COLOR = '0D6EFD'
DATE_COLUMNS = {'Visit Date', 'Next Follow-up'}
DATETIME_COLUMNS = {'Submitted At', 'Reviewed At', 'Created At'}


def _local_naive(dt):
    """Excel/openpyxl cannot store tz-aware datetimes - convert to naive
    local time so real Excel date cells (not text) can be written."""
    if dt is None:
        return None
    return timezone.localtime(dt).replace(tzinfo=None)


EXCEL_COLUMNS = [
    ('DSR Number', 'dsr_number'),
    ('Employee ID', lambda d: d.employee.employee_id),
    ('Employee Name', lambda d: d.employee.get_full_name()),
    ('Visit Date', 'visit_date'),
    ('Client Name', 'client_name'),
    ('Company', 'company_name'),
    ('Contact Person', 'contact_person'),
    ('Contact Number', 'contact_number'),
    ('Project Name', 'project_name'),
    ('Project Type', lambda d: d.get_project_type_display()),
    ('Building Size', lambda d: d.get_building_size_display()),
    ('Project Stage', lambda d: d.get_project_stage_display()),
    ('Purpose of Visit', lambda d: d.get_purpose_of_visit_display()),
    ('Status', lambda d: d.get_status_display()),
    ('Work Done', 'work_done'),
    ('Remarks', 'remarks'),
    ('Next Follow-up', 'next_followup_date'),
    ('Submitted At', lambda d: _local_naive(d.submitted_at)),
    ('Reviewed By', lambda d: d.reviewed_by.get_full_name() if d.reviewed_by else ''),
    ('Reviewed At', lambda d: _local_naive(d.reviewed_at)),
    ('Admin Remarks', 'admin_remarks'),
    ('Created At', lambda d: _local_naive(d.created_at)),
]


def build_dsr_excel_workbook(queryset):
    """Builds the whole workbook in memory. Fine at current/expected scale
    (dozens to low-thousands of rows); if row counts grow much larger,
    Workbook(write_only=True) should replace this to stream rows instead of
    holding the entire sheet in memory."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'DSR Reports'

    header_fill = PatternFill(start_color=HEADER_FILL_COLOR, end_color=HEADER_FILL_COLOR, fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF')

    ws.append([label for label, _ in EXCEL_COLUMNS])
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical='center')
    ws.freeze_panes = 'A2'

    for dsr in queryset:
        row = []
        for _, accessor in EXCEL_COLUMNS:
            value = accessor(dsr) if callable(accessor) else getattr(dsr, accessor)
            row.append(value)
        ws.append(row)

    last_row = ws.max_row
    for idx, (label, _) in enumerate(EXCEL_COLUMNS, start=1):
        if label in DATE_COLUMNS:
            number_format = 'YYYY-MM-DD'
        elif label in DATETIME_COLUMNS:
            number_format = 'YYYY-MM-DD HH:MM'
        else:
            continue
        for row_idx in range(2, last_row + 1):
            ws.cell(row=row_idx, column=idx).number_format = number_format

    # Auto-ish column widths from header + a capped sample of row content,
    # so width computation stays cheap even on a large export.
    sample_row_limit = min(last_row, 200)
    for idx, (label, _) in enumerate(EXCEL_COLUMNS, start=1):
        col_letter = get_column_letter(idx)
        max_len = len(label)
        for row_idx in range(2, sample_row_limit + 1):
            value = ws.cell(row=row_idx, column=idx).value
            if value is not None:
                max_len = max(max_len, len(str(value)))
        ws.column_dimensions[col_letter].width = min(max_len + 2, 40)

    return wb
