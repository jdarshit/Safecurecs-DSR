import io
import os
import shutil
import tempfile
import threading
from datetime import date, timedelta

import openpyxl
from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection
from django.test import TestCase, TransactionTestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from PIL import Image

from .exceptions import InvalidStatusTransition
from .models import DSR, DSRAttachment, DSRDailySequence
from .utils import generate_dsr_number

User = get_user_model()


class TempMediaRootMixin:
    """Isolates file uploads to a throwaway MEDIA_ROOT. Django's test runner
    swaps the database for a test DB automatically, but not file storage, so
    without this every run would leave real files behind under media/."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls._media_root = tempfile.mkdtemp()
        cls._media_root_override = override_settings(MEDIA_ROOT=cls._media_root)
        cls._media_root_override.enable()

    @classmethod
    def tearDownClass(cls):
        cls._media_root_override.disable()
        shutil.rmtree(cls._media_root, ignore_errors=True)
        super().tearDownClass()


def valid_dsr_post_data(**overrides):
    data = {
        'visit_date': '2026-07-06',
        'purpose_of_visit': DSR.PurposeOfVisit.SERVICE,
        'client_name': 'Acme Corp',
        'company_name': 'Acme Corp Pvt Ltd',
        'contact_person': 'John Client',
        'contact_number': '9999999999',
        'client_email': '',
        'project_name': 'Acme HQ Fire System',
        'project_address': '123 Acme Street',
        'google_maps_link': '',
        'project_type': DSR.ProjectType.COMMERCIAL,
        'building_size': DSR.BuildingSize.MEDIUM,
        'built_up_area': '',
        'project_stage': DSR.ProjectStage.EXISTING_BUILDING,
        'work_done': 'Inspected fire panel and sensors.',
        'remarks': '',
        'next_followup_date': '',
    }
    data.update(overrides)
    return data


def make_employee(employee_id='EMP100', **kwargs):
    defaults = dict(email=f'{employee_id.lower()}@example.com', password='TestPass123',
                     first_name='Test', last_name='Employee')
    defaults.update(kwargs)
    return User.objects.create_user(employee_id=employee_id, **defaults)


def make_admin(employee_id='ADM100', **kwargs):
    defaults = dict(email=f'{employee_id.lower()}@example.com', password='AdminPass123',
                     first_name='Test', last_name='Admin', role=User.Role.ADMIN)
    defaults.update(kwargs)
    return User.objects.create_user(employee_id=employee_id, **defaults)


def make_dsr(employee, **kwargs):
    defaults = dict(
        visit_date=date.today(),
        purpose_of_visit=DSR.PurposeOfVisit.SERVICE,
        client_name='Acme Corp', company_name='Acme Corp Pvt Ltd',
        contact_person='John Client', contact_number='9999999999',
        project_name='Acme HQ Fire System', project_address='123 Acme Street',
        project_type=DSR.ProjectType.COMMERCIAL, building_size=DSR.BuildingSize.MEDIUM,
        project_stage=DSR.ProjectStage.EXISTING_BUILDING,
        work_done='Inspected fire panel and sensors.',
    )
    defaults.update(kwargs)
    return DSR.objects.create(employee=employee, **defaults)


def make_jpeg_bytes():
    buf = io.BytesIO()
    Image.new('RGB', (10, 10), color='red').save(buf, format='JPEG')
    return buf.getvalue()


def make_png_bytes():
    buf = io.BytesIO()
    Image.new('RGB', (10, 10), color='blue').save(buf, format='PNG')
    return buf.getvalue()


class DSRNumberGenerationTests(TestCase):
    def setUp(self):
        self.employee = make_employee()

    def test_format(self):
        dsr = make_dsr(self.employee)
        self.assertRegex(dsr.dsr_number, r'^DSR-\d{8}-\d{4}$')

    def test_sequential_same_day(self):
        numbers = [make_dsr(self.employee).dsr_number for _ in range(3)]
        suffixes = [n.split('-')[-1] for n in numbers]
        self.assertEqual(suffixes, ['0001', '0002', '0003'])

    def test_resets_next_day(self):
        today = timezone.localdate()
        tomorrow = today + timedelta(days=1)

        first = make_dsr(self.employee)
        self.assertTrue(first.dsr_number.endswith('-0001'))

        original_localdate = timezone.localdate
        timezone.localdate = lambda: tomorrow
        try:
            second = make_dsr(self.employee)
        finally:
            timezone.localdate = original_localdate

        self.assertTrue(second.dsr_number.endswith('-0001'))
        self.assertIn(tomorrow.strftime('%Y%m%d'), second.dsr_number)
        self.assertEqual(DSRDailySequence.objects.count(), 2)


class DSRNumberConcurrencyTests(TransactionTestCase):
    """Uses TransactionTestCase (not TestCase) because TestCase wraps each test in
    an outer transaction invisible to other threads/connections, which would make
    select_for_update() from a second thread block forever on a lock nothing will
    release. Here each thread gets its own real connection and really commits."""

    def test_concurrent_creation_is_race_safe(self):
        employee = make_employee()
        thread_count = 8
        barrier = threading.Barrier(thread_count)
        results = []
        errors = []
        lock = threading.Lock()

        def worker():
            try:
                barrier.wait(timeout=10)
                number = generate_dsr_number()
                with lock:
                    results.append(number)
            except Exception as exc:  # noqa: BLE001
                with lock:
                    errors.append(exc)
            finally:
                connection.close()

        threads = [threading.Thread(target=worker) for _ in range(thread_count)]
        for t in threads:
            t.start()
        for t in threads:
            t.join(timeout=15)

        self.assertEqual(errors, [])
        self.assertEqual(len(results), thread_count)
        self.assertEqual(len(set(results)), thread_count, 'expected all DSR numbers to be unique')

        suffixes = sorted(int(n.split('-')[-1]) for n in results)
        self.assertEqual(suffixes, list(range(1, thread_count + 1)))


class DSRStatusTransitionTests(TestCase):
    def setUp(self):
        self.employee = make_employee()
        self.admin = make_admin()

    def test_submit_from_draft(self):
        dsr = make_dsr(self.employee)
        dsr.submit()
        dsr.refresh_from_db()
        self.assertEqual(dsr.status, DSR.Status.SUBMITTED)
        self.assertIsNotNone(dsr.submitted_at)

    def test_submit_from_rejected(self):
        dsr = make_dsr(self.employee)
        dsr.submit()
        dsr.reject(self.admin, 'Please add more detail.')
        dsr.submit()
        dsr.refresh_from_db()
        self.assertEqual(dsr.status, DSR.Status.SUBMITTED)

    def test_submit_without_work_done_raises(self):
        dsr = make_dsr(self.employee, work_done='')
        with self.assertRaises(ValidationError):
            dsr.submit()

    def test_approve(self):
        dsr = make_dsr(self.employee)
        dsr.submit()
        dsr.approve(self.admin)
        dsr.refresh_from_db()
        self.assertEqual(dsr.status, DSR.Status.APPROVED)
        self.assertEqual(dsr.reviewed_by_id, self.admin.pk)
        self.assertIsNotNone(dsr.reviewed_at)

    def test_reject_requires_admin_remarks(self):
        dsr = make_dsr(self.employee)
        dsr.submit()
        with self.assertRaises(ValidationError):
            dsr.reject(self.admin, '')
        with self.assertRaises(ValidationError):
            dsr.reject(self.admin, '   ')

    def test_reject_with_remarks(self):
        dsr = make_dsr(self.employee)
        dsr.submit()
        dsr.reject(self.admin, 'Missing client signature.')
        dsr.refresh_from_db()
        self.assertEqual(dsr.status, DSR.Status.REJECTED)
        self.assertEqual(dsr.admin_remarks, 'Missing client signature.')
        self.assertEqual(dsr.reviewed_by_id, self.admin.pk)
        self.assertIsNotNone(dsr.reviewed_at)

    def test_send_back(self):
        dsr = make_dsr(self.employee)
        dsr.submit()
        dsr.send_back(self.admin, 'Please recheck project address.')
        dsr.refresh_from_db()
        self.assertEqual(dsr.status, DSR.Status.DRAFT)
        self.assertEqual(dsr.admin_remarks, 'Please recheck project address.')

    def test_invalid_transitions_raise(self):
        cases = []

        draft = make_dsr(self.employee)
        cases.append((draft, 'approve', (self.admin,)))
        cases.append((draft, 'reject', (self.admin, 'no')))
        cases.append((draft, 'send_back', (self.admin,)))

        submitted = make_dsr(self.employee)
        submitted.submit()
        cases.append((submitted, 'submit', ()))

        rejected = make_dsr(self.employee)
        rejected.submit()
        rejected.reject(self.admin, 'fix this')
        cases.append((rejected, 'approve', (self.admin,)))
        cases.append((rejected, 'reject', (self.admin, 'no')))
        cases.append((rejected, 'send_back', (self.admin,)))

        approved = make_dsr(self.employee)
        approved.submit()
        approved.approve(self.admin)
        cases.append((approved, 'submit', ()))
        cases.append((approved, 'approve', (self.admin,)))
        cases.append((approved, 'reject', (self.admin, 'no')))
        cases.append((approved, 'send_back', (self.admin,)))

        for dsr, method_name, args in cases:
            with self.assertRaises(InvalidStatusTransition, msg=f'{dsr.status} -> {method_name}'):
                getattr(dsr, method_name)(*args)


class DSRCanEditTests(TestCase):
    def setUp(self):
        self.owner = make_employee('EMP101')
        self.other_employee = make_employee('EMP102')
        self.admin = make_admin()

    def test_owner_can_edit_draft_and_rejected(self):
        draft = make_dsr(self.owner)
        self.assertTrue(draft.can_edit(self.owner))

        rejected = make_dsr(self.owner)
        rejected.submit()
        rejected.reject(self.admin, 'fix it')
        self.assertTrue(rejected.can_edit(self.owner))

    def test_owner_cannot_edit_submitted_or_approved(self):
        submitted = make_dsr(self.owner)
        submitted.submit()
        self.assertFalse(submitted.can_edit(self.owner))

        approved = make_dsr(self.owner)
        approved.submit()
        approved.approve(self.admin)
        self.assertFalse(approved.can_edit(self.owner))

    def test_non_owner_can_never_edit(self):
        for status_setup in ('draft', 'submitted', 'rejected', 'approved'):
            dsr = make_dsr(self.owner)
            if status_setup in ('submitted', 'rejected', 'approved'):
                dsr.submit()
            if status_setup == 'rejected':
                dsr.reject(self.admin, 'fix it')
            if status_setup == 'approved':
                dsr.approve(self.admin)
            self.assertFalse(dsr.can_edit(self.other_employee), msg=status_setup)
            self.assertFalse(dsr.can_edit(self.admin), msg=status_setup)


class DSRAttachmentValidationTests(TempMediaRootMixin, TestCase):
    def setUp(self):
        self.employee = make_employee()
        self.dsr = make_dsr(self.employee)

    def test_valid_jpeg_accepted(self):
        upload = SimpleUploadedFile('photo.jpg', make_jpeg_bytes(), content_type='image/jpeg')
        attachment = DSRAttachment(dsr=self.dsr, file=upload, original_filename='photo.jpg')
        attachment.full_clean()
        attachment.save()
        self.assertTrue(attachment.file.name.startswith(f'dsr_attachments/{self.dsr.dsr_number}/'))

    def test_valid_png_accepted(self):
        upload = SimpleUploadedFile('photo.png', make_png_bytes(), content_type='image/png')
        attachment = DSRAttachment(dsr=self.dsr, file=upload, original_filename='photo.png')
        attachment.full_clean()
        attachment.save()

    def test_valid_pdf_accepted(self):
        upload = SimpleUploadedFile('report.pdf', b'%PDF-1.4\n%mock pdf content\n', content_type='application/pdf')
        attachment = DSRAttachment(dsr=self.dsr, file=upload, original_filename='report.pdf')
        attachment.full_clean()
        attachment.save()

    def test_disallowed_extension_rejected(self):
        for name, content_type in [('malware.exe', 'application/octet-stream'), ('notes.txt', 'text/plain')]:
            upload = SimpleUploadedFile(name, b'irrelevant content', content_type=content_type)
            attachment = DSRAttachment(dsr=self.dsr, file=upload, original_filename=name)
            with self.assertRaises(ValidationError, msg=name):
                attachment.full_clean()

    @override_settings(MAX_DSR_ATTACHMENT_SIZE_MB=1)
    def test_oversize_file_rejected(self):
        oversized = make_jpeg_bytes() + (b'0' * (1024 * 1024 + 1))
        upload = SimpleUploadedFile('big.jpg', oversized, content_type='image/jpeg')
        attachment = DSRAttachment(dsr=self.dsr, file=upload, original_filename='big.jpg')
        with self.assertRaises(ValidationError):
            attachment.full_clean()

    def test_content_type_spoofing_rejected(self):
        # Plain text disguised as an image.
        upload = SimpleUploadedFile('fake.jpg', b'this is not an image, just text', content_type='image/jpeg')
        attachment = DSRAttachment(dsr=self.dsr, file=upload, original_filename='fake.jpg')
        with self.assertRaises(ValidationError):
            attachment.full_clean()

    def test_pdf_missing_magic_bytes_rejected(self):
        upload = SimpleUploadedFile('fake.pdf', b'not really a pdf file', content_type='application/pdf')
        attachment = DSRAttachment(dsr=self.dsr, file=upload, original_filename='fake.pdf')
        with self.assertRaises(ValidationError):
            attachment.full_clean()

    def test_upload_path_uses_dsr_number(self):
        upload = SimpleUploadedFile('photo.jpg', make_jpeg_bytes(), content_type='image/jpeg')
        attachment = DSRAttachment.objects.create(dsr=self.dsr, file=upload, original_filename='photo.jpg')
        self.assertTrue(attachment.file.name.startswith(f'dsr_attachments/{self.dsr.dsr_number}/'))


class DSRViewOwnershipTests(TempMediaRootMixin, TestCase):
    def setUp(self):
        self.owner = make_employee('EMP200')
        self.other = make_employee('EMP201')
        self.admin = make_admin('ADM200')
        self.dsr = make_dsr(self.owner)

    def test_other_employee_cannot_view_detail(self):
        self.client.login(username='EMP201', password='TestPass123')
        response = self.client.get(reverse('dsr:detail', args=[self.dsr.pk]))
        self.assertEqual(response.status_code, 404)

    def test_other_employee_cannot_edit(self):
        self.client.login(username='EMP201', password='TestPass123')
        response = self.client.get(reverse('dsr:edit', args=[self.dsr.pk]))
        self.assertEqual(response.status_code, 404)

    def test_other_employee_cannot_delete_attachment(self):
        attachment = DSRAttachment.objects.create(
            dsr=self.dsr, file=SimpleUploadedFile('a.jpg', make_jpeg_bytes(), content_type='image/jpeg'),
            original_filename='a.jpg', category=DSRAttachment.Category.ATTACHMENT,
        )
        self.client.login(username='EMP201', password='TestPass123')
        response = self.client.post(reverse('dsr:attachment_delete', args=[self.dsr.pk, attachment.pk]))
        self.assertEqual(response.status_code, 404)
        self.assertTrue(DSRAttachment.objects.filter(pk=attachment.pk).exists())

    def test_admin_blocked_from_create(self):
        self.client.login(username='ADM200', password='AdminPass123')
        response = self.client.get(reverse('dsr:create'))
        self.assertRedirects(response, reverse('accounts:admin_dashboard'))

    def test_admin_blocked_from_my_reports(self):
        self.client.login(username='ADM200', password='AdminPass123')
        response = self.client.get(reverse('dsr:my_reports'))
        self.assertRedirects(response, reverse('accounts:admin_dashboard'))


class DSRCreateViewTests(TempMediaRootMixin, TestCase):
    def setUp(self):
        self.employee = make_employee('EMP210')
        self.client.login(username='EMP210', password='TestPass123')

    def test_save_draft_minimal_fields(self):
        data = {'action': 'save_draft', 'visit_date': '2026-07-06', 'client_name': 'Acme'}
        response = self.client.post(reverse('dsr:create'), data)
        self.assertEqual(DSR.objects.count(), 1)
        dsr = DSR.objects.get()
        self.assertEqual(dsr.status, DSR.Status.DRAFT)
        self.assertEqual(dsr.employee, self.employee)
        self.assertRedirects(response, reverse('dsr:detail', args=[dsr.pk]))

    def test_submit_missing_required_fields_shows_errors_and_does_not_create(self):
        data = {'action': 'submit', 'visit_date': '2026-07-06', 'client_name': 'Acme'}
        response = self.client.post(reverse('dsr:create'), data)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(DSR.objects.count(), 0)

    def test_submit_complete_data_transitions_to_submitted(self):
        data = valid_dsr_post_data(action='submit')
        response = self.client.post(reverse('dsr:create'), data)
        dsr = DSR.objects.get()
        self.assertEqual(dsr.status, DSR.Status.SUBMITTED)
        self.assertIsNotNone(dsr.submitted_at)
        self.assertRedirects(response, reverse('dsr:detail', args=[dsr.pk]))

    def test_tampering_employee_status_dsr_number_ignored(self):
        other = make_employee('EMP211')
        data = valid_dsr_post_data(action='save_draft')
        data['employee'] = other.pk
        data['status'] = DSR.Status.APPROVED
        data['dsr_number'] = 'HACKED-0000'
        self.client.post(reverse('dsr:create'), data)
        dsr = DSR.objects.get()
        self.assertEqual(dsr.employee, self.employee)
        self.assertEqual(dsr.status, DSR.Status.DRAFT)
        self.assertNotEqual(dsr.dsr_number, 'HACKED-0000')


class DSRUpdateViewTests(TempMediaRootMixin, TestCase):
    def setUp(self):
        self.employee = make_employee('EMP220')
        self.admin = make_admin('ADM220')
        self.client.login(username='EMP220', password='TestPass123')

    def test_edit_blocked_when_submitted(self):
        dsr = make_dsr(self.employee)
        dsr.submit()
        response = self.client.get(reverse('dsr:edit', args=[dsr.pk]))
        self.assertRedirects(response, reverse('dsr:detail', args=[dsr.pk]))

    def test_edit_blocked_when_approved(self):
        dsr = make_dsr(self.employee)
        dsr.submit()
        dsr.approve(self.admin)
        response = self.client.get(reverse('dsr:edit', args=[dsr.pk]))
        self.assertRedirects(response, reverse('dsr:detail', args=[dsr.pk]))

    def test_edit_allowed_when_draft(self):
        dsr = make_dsr(self.employee)
        response = self.client.get(reverse('dsr:edit', args=[dsr.pk]))
        self.assertEqual(response.status_code, 200)

    def test_edit_allowed_when_rejected(self):
        dsr = make_dsr(self.employee)
        dsr.submit()
        dsr.reject(self.admin, 'fix it')
        response = self.client.get(reverse('dsr:edit', args=[dsr.pk]))
        self.assertEqual(response.status_code, 200)

    def test_rejected_resubmit_transitions_to_submitted(self):
        dsr = make_dsr(self.employee)
        dsr.submit()
        dsr.reject(self.admin, 'fix it')
        data = valid_dsr_post_data(action='submit')
        response = self.client.post(reverse('dsr:edit', args=[dsr.pk]), data)
        dsr.refresh_from_db()
        self.assertEqual(dsr.status, DSR.Status.SUBMITTED)
        self.assertRedirects(response, reverse('dsr:detail', args=[dsr.pk]))

    def test_save_draft_on_rejected_keeps_rejected_status_and_remarks(self):
        dsr = make_dsr(self.employee)
        dsr.submit()
        dsr.reject(self.admin, 'Please fix the address.')
        data = valid_dsr_post_data(action='save_draft', project_address='Updated Address')
        response = self.client.post(reverse('dsr:edit', args=[dsr.pk]), data)
        dsr.refresh_from_db()
        self.assertEqual(dsr.status, DSR.Status.REJECTED)
        self.assertEqual(dsr.admin_remarks, 'Please fix the address.')
        self.assertEqual(dsr.project_address, 'Updated Address')
        self.assertRedirects(response, reverse('dsr:detail', args=[dsr.pk]))

    def test_tampering_on_edit_ignored(self):
        dsr = make_dsr(self.employee)
        other = make_employee('EMP221')
        data = valid_dsr_post_data(action='save_draft')
        data['employee'] = other.pk
        data['status'] = DSR.Status.APPROVED
        data['dsr_number'] = 'HACKED'
        self.client.post(reverse('dsr:edit', args=[dsr.pk]), data)
        dsr.refresh_from_db()
        self.assertEqual(dsr.employee, self.employee)
        self.assertEqual(dsr.status, DSR.Status.DRAFT)
        self.assertNotEqual(dsr.dsr_number, 'HACKED')


class DSRAttachmentUploadViewTests(TempMediaRootMixin, TestCase):
    def setUp(self):
        self.employee = make_employee('EMP230')
        self.client.login(username='EMP230', password='TestPass123')

    def test_multiple_files_attach_to_correct_categories(self):
        data = valid_dsr_post_data(action='save_draft')
        data['project_photos'] = [
            SimpleUploadedFile('p1.jpg', make_jpeg_bytes(), content_type='image/jpeg'),
            SimpleUploadedFile('p2.png', make_png_bytes(), content_type='image/png'),
        ]
        data['attachments'] = [
            SimpleUploadedFile('doc.pdf', b'%PDF-1.4\n%mock pdf content\n', content_type='application/pdf'),
        ]
        self.client.post(reverse('dsr:create'), data)
        dsr = DSR.objects.get()
        self.assertEqual(dsr.attachments.filter(category=DSRAttachment.Category.PROJECT_PHOTO).count(), 2)
        self.assertEqual(dsr.attachments.filter(category=DSRAttachment.Category.ATTACHMENT).count(), 1)

    def test_invalid_file_type_rejected_nothing_saved(self):
        data = valid_dsr_post_data(action='save_draft')
        data['attachments'] = [SimpleUploadedFile('bad.exe', b'irrelevant', content_type='application/octet-stream')]
        response = self.client.post(reverse('dsr:create'), data)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(DSR.objects.count(), 0)
        self.assertEqual(DSRAttachment.objects.count(), 0)

    @override_settings(MAX_DSR_ATTACHMENT_SIZE_MB=1)
    def test_oversize_file_rejected_nothing_saved(self):
        oversized = make_jpeg_bytes() + (b'0' * (1024 * 1024 + 1))
        data = valid_dsr_post_data(action='save_draft')
        data['project_photos'] = [SimpleUploadedFile('big.jpg', oversized, content_type='image/jpeg')]
        response = self.client.post(reverse('dsr:create'), data)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(DSR.objects.count(), 0)

    def test_attachment_delete_by_owner_when_editable(self):
        dsr = make_dsr(self.employee)
        attachment = DSRAttachment.objects.create(
            dsr=dsr, file=SimpleUploadedFile('a.jpg', make_jpeg_bytes(), content_type='image/jpeg'),
            original_filename='a.jpg', category=DSRAttachment.Category.ATTACHMENT,
        )
        response = self.client.post(reverse('dsr:attachment_delete', args=[dsr.pk, attachment.pk]))
        self.assertRedirects(response, reverse('dsr:edit', args=[dsr.pk]))
        self.assertFalse(DSRAttachment.objects.filter(pk=attachment.pk).exists())

    def test_attachment_delete_blocked_when_not_editable(self):
        dsr = make_dsr(self.employee)
        dsr.submit()
        attachment = DSRAttachment.objects.create(
            dsr=dsr, file=SimpleUploadedFile('a.jpg', make_jpeg_bytes(), content_type='image/jpeg'),
            original_filename='a.jpg', category=DSRAttachment.Category.ATTACHMENT,
        )
        response = self.client.post(reverse('dsr:attachment_delete', args=[dsr.pk, attachment.pk]))
        self.assertRedirects(response, reverse('dsr:detail', args=[dsr.pk]))
        self.assertTrue(DSRAttachment.objects.filter(pk=attachment.pk).exists())


class DSRListViewTests(TestCase):
    def setUp(self):
        self.employee = make_employee('EMP240')
        self.other = make_employee('EMP241')
        self.client.login(username='EMP240', password='TestPass123')

    def test_only_own_reports_listed(self):
        mine = make_dsr(self.employee)
        make_dsr(self.other)
        response = self.client.get(reverse('dsr:my_reports'))
        self.assertEqual(list(response.context['dsrs']), [mine])

    def test_status_filter(self):
        make_dsr(self.employee)
        submitted = make_dsr(self.employee)
        submitted.submit()
        response = self.client.get(reverse('dsr:my_reports'), {'status': DSR.Status.SUBMITTED})
        self.assertEqual(list(response.context['dsrs']), [submitted])

    def test_date_range_filter(self):
        make_dsr(self.employee, visit_date=date(2026, 1, 1))
        recent = make_dsr(self.employee, visit_date=date(2026, 7, 1))
        response = self.client.get(reverse('dsr:my_reports'), {'date_from': '2026-06-01'})
        self.assertEqual(list(response.context['dsrs']), [recent])

    def test_search(self):
        match = make_dsr(self.employee, client_name='Zephyr Industries')
        make_dsr(self.employee, client_name='Other Client')
        response = self.client.get(reverse('dsr:my_reports'), {'q': 'zephyr'})
        self.assertEqual(list(response.context['dsrs']), [match])

    def test_pagination(self):
        for i in range(17):
            make_dsr(self.employee, visit_date=date(2026, 1, 1) + timedelta(days=i))
        response = self.client.get(reverse('dsr:my_reports'))
        self.assertEqual(len(response.context['dsrs']), 15)
        response_page2 = self.client.get(reverse('dsr:my_reports'), {'page': 2})
        self.assertEqual(len(response_page2.context['dsrs']), 2)

    def test_empty_state(self):
        response = self.client.get(reverse('dsr:my_reports'))
        self.assertContains(response, 'No reports found')


class DSRSendBackModelValidationTests(TestCase):
    def test_send_back_without_remarks_raises_and_leaves_unchanged(self):
        employee = make_employee('EMP300')
        admin = make_admin('ADM300')
        dsr = make_dsr(employee)
        dsr.submit()
        with self.assertRaises(ValidationError):
            dsr.send_back(admin, admin_remarks='')
        dsr.refresh_from_db()
        self.assertEqual(dsr.status, DSR.Status.SUBMITTED)
        self.assertEqual(dsr.admin_remarks, '')


class AdminDSRAccessTests(TestCase):
    def setUp(self):
        self.employee = make_employee('EMP310')
        self.admin = make_admin('ADM310')
        self.dsr = make_dsr(self.employee)
        self.dsr.submit()

    def test_employee_blocked_from_all_reports(self):
        self.client.login(username='EMP310', password='TestPass123')
        response = self.client.get(reverse('dsr:all_reports'))
        self.assertRedirects(response, reverse('accounts:employee_dashboard'))

    def test_employee_blocked_from_review(self):
        self.client.login(username='EMP310', password='TestPass123')
        response = self.client.get(reverse('dsr:review', args=[self.dsr.pk]))
        self.assertRedirects(response, reverse('accounts:employee_dashboard'))

    def test_admin_can_access_all_reports(self):
        self.client.login(username='ADM310', password='AdminPass123')
        response = self.client.get(reverse('dsr:all_reports'))
        self.assertEqual(response.status_code, 200)

    def test_admin_can_access_review(self):
        self.client.login(username='ADM310', password='AdminPass123')
        response = self.client.get(reverse('dsr:review', args=[self.dsr.pk]))
        self.assertEqual(response.status_code, 200)


class AdminDSRListViewTests(TestCase):
    def setUp(self):
        self.admin = make_admin('ADM320')
        self.emp1 = make_employee('EMP320', first_name='Alice', last_name='Anderson')
        self.emp2 = make_employee('EMP321', first_name='Bob', last_name='Brown')
        self.client.login(username='ADM320', password='AdminPass123')

    def test_shows_reports_from_multiple_employees(self):
        d1 = make_dsr(self.emp1)
        d2 = make_dsr(self.emp2)
        response = self.client.get(reverse('dsr:all_reports'))
        ids = {d.pk for d in response.context['dsrs']}
        self.assertEqual(ids, {d1.pk, d2.pk})

    def test_employee_filter(self):
        mine = make_dsr(self.emp1)
        make_dsr(self.emp2)
        response = self.client.get(reverse('dsr:all_reports'), {'employee': self.emp1.pk})
        self.assertEqual(list(response.context['dsrs']), [mine])

    def test_client_filter(self):
        match = make_dsr(self.emp1, client_name='Zephyr Industries')
        make_dsr(self.emp1, client_name='Other Client')
        response = self.client.get(reverse('dsr:all_reports'), {'client': 'zephyr'})
        self.assertEqual(list(response.context['dsrs']), [match])

    def test_project_type_filter(self):
        match = make_dsr(self.emp1, project_type=DSR.ProjectType.HOSPITAL)
        make_dsr(self.emp1, project_type=DSR.ProjectType.OFFICE)
        response = self.client.get(reverse('dsr:all_reports'), {'project_type': DSR.ProjectType.HOSPITAL})
        self.assertEqual(list(response.context['dsrs']), [match])

    def test_status_filter(self):
        make_dsr(self.emp1)
        submitted = make_dsr(self.emp1)
        submitted.submit()
        response = self.client.get(reverse('dsr:all_reports'), {'status': DSR.Status.SUBMITTED})
        self.assertEqual(list(response.context['dsrs']), [submitted])

    def test_date_range_filter(self):
        make_dsr(self.emp1, visit_date=date(2026, 1, 1))
        recent = make_dsr(self.emp1, visit_date=date(2026, 7, 1))
        response = self.client.get(reverse('dsr:all_reports'), {'date_from': '2026-06-01'})
        self.assertEqual(list(response.context['dsrs']), [recent])

    def test_search_by_employee_name(self):
        match = make_dsr(self.emp1)
        make_dsr(self.emp2)
        response = self.client.get(reverse('dsr:all_reports'), {'q': 'Alice'})
        self.assertEqual(list(response.context['dsrs']), [match])

    def test_search_by_employee_id(self):
        match = make_dsr(self.emp1)
        make_dsr(self.emp2)
        response = self.client.get(reverse('dsr:all_reports'), {'q': 'EMP320'})
        self.assertEqual(list(response.context['dsrs']), [match])

    def test_combined_filters(self):
        match = make_dsr(self.emp1, client_name='Zephyr Industries', project_type=DSR.ProjectType.HOSPITAL)
        make_dsr(self.emp1, client_name='Zephyr Industries', project_type=DSR.ProjectType.OFFICE)
        make_dsr(self.emp2, client_name='Zephyr Industries', project_type=DSR.ProjectType.HOSPITAL)
        response = self.client.get(reverse('dsr:all_reports'), {
            'employee': self.emp1.pk, 'client': 'zephyr', 'project_type': DSR.ProjectType.HOSPITAL,
        })
        self.assertEqual(list(response.context['dsrs']), [match])

    def test_pagination(self):
        for i in range(22):
            make_dsr(self.emp1, visit_date=date(2026, 1, 1) + timedelta(days=i))
        response = self.client.get(reverse('dsr:all_reports'))
        self.assertEqual(len(response.context['dsrs']), 20)
        response_page2 = self.client.get(reverse('dsr:all_reports'), {'page': 2})
        self.assertEqual(len(response_page2.context['dsrs']), 2)

    def test_status_tab_counts(self):
        make_dsr(self.emp1)
        submitted = make_dsr(self.emp1)
        submitted.submit()
        response = self.client.get(reverse('dsr:all_reports'))
        tabs = {tab['value']: tab['count'] for tab in response.context['status_tabs']}
        self.assertEqual(tabs[''], 2)
        self.assertEqual(tabs[DSR.Status.SUBMITTED], 1)
        self.assertEqual(tabs[DSR.Status.DRAFT], 1)
        self.assertEqual(tabs[DSR.Status.APPROVED], 0)
        self.assertEqual(tabs[DSR.Status.REJECTED], 0)


class AdminDSRReviewViewTests(TestCase):
    def setUp(self):
        self.admin = make_admin('ADM330')
        self.employee = make_employee('EMP330')
        self.client.login(username='ADM330', password='AdminPass123')

    def _submitted_dsr(self):
        dsr = make_dsr(self.employee)
        dsr.submit()
        return dsr

    def test_approve_sets_fields(self):
        dsr = self._submitted_dsr()
        response = self.client.post(
            reverse('dsr:review', args=[dsr.pk]), {'action': 'approve', 'admin_remarks': ''},
        )
        dsr.refresh_from_db()
        self.assertEqual(dsr.status, DSR.Status.APPROVED)
        self.assertEqual(dsr.reviewed_by, self.admin)
        self.assertIsNotNone(dsr.reviewed_at)
        self.assertRedirects(response, reverse('dsr:review', args=[dsr.pk]))

    def test_reject_without_remarks_fails_with_error(self):
        dsr = self._submitted_dsr()
        response = self.client.post(
            reverse('dsr:review', args=[dsr.pk]), {'action': 'reject', 'admin_remarks': ''},
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Admin remarks are required')
        dsr.refresh_from_db()
        self.assertEqual(dsr.status, DSR.Status.SUBMITTED)

    def test_reject_with_remarks_works(self):
        dsr = self._submitted_dsr()
        response = self.client.post(
            reverse('dsr:review', args=[dsr.pk]), {'action': 'reject', 'admin_remarks': 'Fix the address'},
        )
        dsr.refresh_from_db()
        self.assertEqual(dsr.status, DSR.Status.REJECTED)
        self.assertEqual(dsr.admin_remarks, 'Fix the address')
        self.assertRedirects(response, reverse('dsr:review', args=[dsr.pk]))

    def test_send_back_without_remarks_fails_with_error(self):
        dsr = self._submitted_dsr()
        response = self.client.post(
            reverse('dsr:review', args=[dsr.pk]), {'action': 'send_back', 'admin_remarks': ''},
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Admin remarks are required')
        dsr.refresh_from_db()
        self.assertEqual(dsr.status, DSR.Status.SUBMITTED)

    def test_send_back_with_remarks_works_and_employee_can_edit(self):
        dsr = self._submitted_dsr()
        response = self.client.post(
            reverse('dsr:review', args=[dsr.pk]), {'action': 'send_back', 'admin_remarks': 'Please recheck'},
        )
        dsr.refresh_from_db()
        self.assertEqual(dsr.status, DSR.Status.DRAFT)
        self.assertEqual(dsr.admin_remarks, 'Please recheck')
        self.assertRedirects(response, reverse('dsr:review', args=[dsr.pk]))

        self.client.logout()
        self.client.login(username='EMP330', password='TestPass123')
        edit_response = self.client.get(reverse('dsr:edit', args=[dsr.pk]))
        self.assertEqual(edit_response.status_code, 200)
        self.assertContains(edit_response, 'Sent Back for Correction')
        self.assertContains(edit_response, 'Please recheck')

    def test_reject_remarks_visible_to_owner_on_detail(self):
        dsr = self._submitted_dsr()
        self.client.post(
            reverse('dsr:review', args=[dsr.pk]), {'action': 'reject', 'admin_remarks': 'Fix the address'},
        )
        self.client.logout()
        self.client.login(username='EMP330', password='TestPass123')
        response = self.client.get(reverse('dsr:detail', args=[dsr.pk]))
        self.assertContains(response, 'Fix the address')

    def test_actions_blocked_on_draft(self):
        dsr = make_dsr(self.employee)
        response = self.client.post(
            reverse('dsr:review', args=[dsr.pk]), {'action': 'approve', 'admin_remarks': ''},
        )
        self.assertRedirects(response, reverse('dsr:review', args=[dsr.pk]))
        dsr.refresh_from_db()
        self.assertEqual(dsr.status, DSR.Status.DRAFT)

    def test_actions_blocked_on_approved(self):
        dsr = self._submitted_dsr()
        dsr.approve(self.admin)
        response = self.client.post(
            reverse('dsr:review', args=[dsr.pk]), {'action': 'reject', 'admin_remarks': 'x'},
        )
        self.assertRedirects(response, reverse('dsr:review', args=[dsr.pk]))
        dsr.refresh_from_db()
        self.assertEqual(dsr.status, DSR.Status.APPROVED)

    def test_get_does_not_mutate_state(self):
        dsr = self._submitted_dsr()
        self.client.get(reverse('dsr:review', args=[dsr.pk]), {'action': 'approve'})
        dsr.refresh_from_db()
        self.assertEqual(dsr.status, DSR.Status.SUBMITTED)


class AllReportsExportAccessTests(TestCase):
    def setUp(self):
        self.employee = make_employee('EMP400')
        self.admin = make_admin('ADM400')
        make_dsr(self.employee)

    def test_employee_blocked_from_excel_export(self):
        self.client.login(username='EMP400', password='TestPass123')
        response = self.client.get(reverse('dsr:all_reports_export_excel'))
        self.assertRedirects(response, reverse('accounts:employee_dashboard'))

    def test_employee_blocked_from_pdf_export(self):
        self.client.login(username='EMP400', password='TestPass123')
        response = self.client.get(reverse('dsr:all_reports_export_pdf'))
        self.assertRedirects(response, reverse('accounts:employee_dashboard'))

    def test_admin_allowed_excel_export(self):
        self.client.login(username='ADM400', password='AdminPass123')
        response = self.client.get(reverse('dsr:all_reports_export_excel'))
        self.assertEqual(response.status_code, 200)

    def test_admin_allowed_pdf_export(self):
        self.client.login(username='ADM400', password='AdminPass123')
        response = self.client.get(reverse('dsr:all_reports_export_pdf'))
        self.assertEqual(response.status_code, 200)


class ExcelExportTests(TestCase):
    def setUp(self):
        self.admin = make_admin('ADM410')
        self.emp1 = make_employee('EMP410', first_name='Uma', last_name='Employee')
        self.emp2 = make_employee('EMP411', first_name='Victor', last_name='Employee')
        self.client.login(username='ADM410', password='AdminPass123')

    def test_content_type_and_filename(self):
        make_dsr(self.emp1)
        response = self.client.get(reverse('dsr:all_reports_export_excel'))
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        self.assertRegex(
            response['Content-Disposition'],
            r'attachment; filename="DSR_Reports_\d{8}-\d{4}\.xlsx"',
        )

    def test_workbook_contents_and_filters_applied(self):
        included = make_dsr(self.emp1, client_name='Zephyr Industries')
        included.submit()
        excluded = make_dsr(self.emp2, client_name='Other Client')  # different employee, must be excluded

        response = self.client.get(
            reverse('dsr:all_reports_export_excel'), {'employee': self.emp1.pk},
        )
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb.active

        headers = [cell.value for cell in ws[1]]
        self.assertEqual(headers[0], 'DSR Number')
        self.assertEqual(headers[1], 'Employee ID')
        self.assertEqual(headers[13], 'Status')

        data_rows = list(ws.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(data_rows), 1)

        row = data_rows[0]
        self.assertEqual(row[0], included.dsr_number)
        self.assertEqual(row[1], self.emp1.employee_id)
        self.assertEqual(row[2], self.emp1.get_full_name())
        self.assertEqual(row[4], 'Zephyr Industries')
        self.assertEqual(row[13], 'Submitted')

        exported_numbers = {r[0] for r in data_rows}
        self.assertNotIn(excluded.dsr_number, exported_numbers)


class ListPDFExportTests(TestCase):
    def setUp(self):
        self.admin = make_admin('ADM420')
        self.employee = make_employee('EMP420')
        self.client.login(username='ADM420', password='AdminPass123')

    def test_pdf_export_smoke(self):
        make_dsr(self.employee)
        response = self.client.get(reverse('dsr:all_reports_export_pdf'))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/pdf')
        self.assertTrue(response.content.startswith(b'%PDF'))
        self.assertRegex(
            response['Content-Disposition'],
            r'attachment; filename="DSR_Reports_\d{8}-\d{4}\.pdf"',
        )


class SharedFilterConsistencyTests(TestCase):
    """Proves the list view and the exports use the same filtering (via
    AdminDSRFilterMixin) rather than parallel, possibly-diverging logic."""

    def setUp(self):
        self.admin = make_admin('ADM430')
        self.emp1 = make_employee('EMP430')
        self.emp2 = make_employee('EMP431')
        self.client.login(username='ADM430', password='AdminPass123')

    def test_list_view_and_excel_export_agree_on_filtered_set(self):
        matching = make_dsr(self.emp1)
        make_dsr(self.emp2)  # must be excluded by the employee filter

        list_response = self.client.get(reverse('dsr:all_reports'), {'employee': self.emp1.pk})
        list_numbers = {d.dsr_number for d in list_response.context['dsrs']}

        excel_response = self.client.get(
            reverse('dsr:all_reports_export_excel'), {'employee': self.emp1.pk},
        )
        wb = openpyxl.load_workbook(io.BytesIO(excel_response.content))
        excel_numbers = {row[0] for row in wb.active.iter_rows(min_row=2, values_only=True)}

        self.assertEqual(list_numbers, {matching.dsr_number})
        self.assertEqual(excel_numbers, {matching.dsr_number})
        self.assertEqual(list_numbers, excel_numbers)


class SingleDSRPDFExportTests(TempMediaRootMixin, TestCase):
    def setUp(self):
        self.admin = make_admin('ADM440')
        self.owner = make_employee('EMP440')
        self.other = make_employee('EMP441')

    def _approved_dsr(self, employee):
        dsr = make_dsr(employee)
        dsr.submit()
        dsr.approve(self.admin)
        return dsr

    def test_owner_can_download_own_approved_pdf(self):
        dsr = self._approved_dsr(self.owner)
        self.client.login(username='EMP440', password='TestPass123')
        response = self.client.get(reverse('dsr:export_pdf', args=[dsr.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/pdf')
        self.assertTrue(response.content.startswith(b'%PDF'))
        self.assertEqual(response['Content-Disposition'], f'attachment; filename="{dsr.dsr_number}.pdf"')

    def test_owner_cannot_download_own_submitted_pdf(self):
        dsr = make_dsr(self.owner)
        dsr.submit()
        self.client.login(username='EMP440', password='TestPass123')
        response = self.client.get(reverse('dsr:export_pdf', args=[dsr.pk]))
        self.assertRedirects(response, reverse('dsr:detail', args=[dsr.pk]))
        dsr.refresh_from_db()
        self.assertEqual(dsr.status, DSR.Status.SUBMITTED)

    def test_employee_cannot_download_another_employees_pdf(self):
        dsr = self._approved_dsr(self.other)
        self.client.login(username='EMP440', password='TestPass123')
        response = self.client.get(reverse('dsr:export_pdf', args=[dsr.pk]))
        self.assertEqual(response.status_code, 404)

    def test_admin_can_download_any_status_pdf(self):
        draft = make_dsr(self.owner)
        self.client.login(username='ADM440', password='AdminPass123')
        response = self.client.get(reverse('dsr:export_pdf', args=[draft.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b'%PDF'))

    def test_missing_image_file_does_not_crash_export(self):
        dsr = self._approved_dsr(self.owner)
        attachment = DSRAttachment.objects.create(
            dsr=dsr,
            file=SimpleUploadedFile('site.jpg', make_jpeg_bytes(), content_type='image/jpeg'),
            original_filename='site.jpg',
            category=DSRAttachment.Category.PROJECT_PHOTO,
        )
        # Delete the file from disk directly, leaving the DB row dangling -
        # simulates a file lost/removed outside the normal delete flow.
        os.remove(attachment.file.path)

        self.client.login(username='ADM440', password='AdminPass123')
        response = self.client.get(reverse('dsr:export_pdf', args=[dsr.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b'%PDF'))


class DSRAttachmentDownloadViewTests(TempMediaRootMixin, TestCase):
    def setUp(self):
        self.admin = make_admin('ADM450')
        self.owner = make_employee('EMP450')
        self.other = make_employee('EMP451')
        self.dsr = make_dsr(self.owner)
        self.attachment = DSRAttachment.objects.create(
            dsr=self.dsr,
            file=SimpleUploadedFile('site.jpg', make_jpeg_bytes(), content_type='image/jpeg'),
            original_filename='site.jpg',
            category=DSRAttachment.Category.PROJECT_PHOTO,
        )

    def test_owner_can_download(self):
        self.client.login(username='EMP450', password='TestPass123')
        response = self.client.get(reverse('dsr:attachment_download', args=[self.attachment.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'image/jpeg')

    def test_admin_can_download_any(self):
        self.client.login(username='ADM450', password='AdminPass123')
        response = self.client.get(reverse('dsr:attachment_download', args=[self.attachment.pk]))
        self.assertEqual(response.status_code, 200)

    def test_other_employee_gets_404(self):
        self.client.login(username='EMP451', password='TestPass123')
        response = self.client.get(reverse('dsr:attachment_download', args=[self.attachment.pk]))
        self.assertEqual(response.status_code, 404)

    def test_anonymous_redirected_to_login(self):
        response = self.client.get(reverse('dsr:attachment_download', args=[self.attachment.pk]))
        expected_url = reverse('dsr:attachment_download', args=[self.attachment.pk])
        self.assertRedirects(response, f"{reverse('accounts:login')}?next={expected_url}")

    def test_raw_media_url_not_served(self):
        # dsr_attachments/ is deliberately NOT mapped by the dev static()
        # helper (only profile_photos/ is) - a raw, guessed URL must 404
        # even for a logged-in user, proving the protection is real, not
        # just "the UI happens not to link there."
        self.client.login(username='EMP450', password='TestPass123')
        raw_url = f'/media/{self.attachment.file.name}'
        response = self.client.get(raw_url)
        self.assertEqual(response.status_code, 404)
